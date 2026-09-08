# -*- coding: utf-8 -*-
"""文档解析 —— 本地迁移时使用，服务器端不跑"""
import os


def parse(path: str) -> tuple[str, bool]:
    """返回 (文本, 是否扫描件)，失败返回 ("", False)"""
    ext = os.path.splitext(path)[1].lower()

    try:
        if ext == ".pdf":
            return _parse_pdf(path)
        elif ext in (".docx",):
            return _parse_docx(path)
        elif ext in (".xlsx",):
            return _parse_xlsx(path)
        elif ext in (".xls",):
            return _parse_xls(path)
        elif ext == ".pptx":
            return _parse_pptx(path)
        elif ext in (".txt", ".md", ".csv", ".rtf"):
            return _parse_text(path)
    except Exception as e:
        print(f"解析失败 {path}: {e}")
    return "", False


def _parse_pdf(path: str) -> tuple[str, bool]:
    import fitz
    import pikepdf
    import tempfile

    # 先用 pikepdf 解密
    tmp_path = path
    try:
        pdf = pikepdf.open(path)
        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".pdf")
        os.close(tmp_fd)
        pdf.save(tmp_path)
        pdf.close()
    except Exception:
        pass

    doc = fitz.open(tmp_path)
    pages = [page.get_text() for page in doc]
    doc.close()

    if tmp_path != path:
        try:
            os.remove(tmp_path)
        except Exception:
            pass

    text = "\n".join(pages).strip()
    is_scanned = len(text) < 50
    return text, is_scanned


def _parse_docx(path: str) -> tuple[str, bool]:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs), False


def _parse_xlsx(path: str) -> tuple[str, bool]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                texts.append("\t".join(cells))
    wb.close()
    return "\n".join(texts), False


def _parse_xls(path: str) -> tuple[str, bool]:
    import xlrd
    wb = xlrd.open_workbook(path)
    texts = []
    for sheet in wb.sheets():
        for r in range(sheet.nrows):
            row = [str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            texts.append("\t".join(row))
    return "\n".join(texts), False


def _parse_pptx(path: str) -> tuple[str, bool]:
    from pptx import Presentation
    prs = Presentation(path)
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append("\n".join(p.text for p in shape.text_frame.paragraphs))
    return "\n".join(texts), False


def _parse_text(path: str) -> tuple[str, bool]:
    for enc in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read(), False
        except UnicodeDecodeError:
            continue
    return "", False
